from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Team_Agent_Design_and_Integration_Standard.docx"

NAVY = "17365D"
PALE_BLUE = "EAF1F8"
PALE_GRAY = "F4F4F4"
LIGHT_GRAY = "D9D9D9"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(89, 89, 89)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=85, start=100, bottom=85, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_layout(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx < len(widths):
                set_cell_width(cell, widths[idx])


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def set_font(run, name="Aptos", size=10.3, bold=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = color


def add_body(doc, text, before=0, after=4.5, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.06
    p.paragraph_format.keep_together = keep
    set_font(p.add_run(text))
    return p


def add_label_paragraph(doc, label, text, before=0, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.04
    r = p.add_run(label)
    set_font(r, bold=True)
    set_font(p.add_run(text))
    return p


def add_bullet(doc, label, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.02
    set_font(p.add_run(label), bold=True)
    set_font(p.add_run(text))
    return p


def add_heading(doc, text, level=1, before=7, after=3):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, size=12.4 if level == 1 else 10.8, bold=True)
    return p


def style_table_text(table, header=True):
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_font(run, size=8.8, bold=(r_idx == 0 and header), color=(RGBColor(255, 255, 255) if r_idx == 0 and header else BLACK))


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.53)
section.bottom_margin = Inches(0.52)
section.left_margin = Inches(0.67)
section.right_margin = Inches(0.67)
section.header_distance = Inches(0.25)
section.footer_distance = Inches(0.25)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.3)
normal.font.color.rgb = BLACK

for style_name in ("Title", "Heading 1", "Heading 2"):
    style = styles[style_name]
    style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
    style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
    style.font.color.rgb = BLACK

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(2)
title.paragraph_format.keep_with_next = True
set_font(title.add_run("Team Agent Design and Integration Standard"), name="Aptos Display", size=18.5, bold=True)
remove_paragraph_borders(title)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(7)
set_font(subtitle.add_run("Team Workshop 2  Emerging Technologies  Agentic AI"), size=9.5, bold=True, color=GRAY)

meta = doc.add_table(rows=2, cols=4)
meta.alignment = WD_TABLE_ALIGNMENT.LEFT
widths = [0.72, 2.75, 0.72, 2.75]
set_table_layout(meta, widths)
meta.cell(0, 0).text = "TEAM"
meta.cell(0, 1).text = "____________________________"
meta.cell(0, 2).text = "DATE"
meta.cell(0, 3).text = "____________________________"
meta.cell(1, 0).text = "MEMBERS"
meta.cell(1, 1).merge(meta.cell(1, 3)).text = "____________________________________________________________________________"
for row in meta.rows:
    shade_cell(row.cells[0], PALE_BLUE)
    shade_cell(row.cells[2], PALE_BLUE)
set_table_borders(meta)
style_table_text(meta, header=False)
for row in meta.rows:
    for idx, cell in enumerate(row.cells):
        for p in cell.paragraphs:
            for run in p.runs:
                set_font(run, size=8.5, bold=(idx in (0, 2)))

add_heading(doc, "Purpose and Team Decision", before=7)
add_body(doc, "This standard governs how our team evaluates independently developed specialist agents, selects agents for a shared system, and preserves evidence for the final project. We will accept an agent only when its behavior is traceable, comparable with other candidates, and compatible with the shared architecture. Passing a file-format check alone is not sufficient.")

add_heading(doc, "Shared Architecture and Editable Scope")
add_bullet(doc, "Frozen core  ", "The shared common instructions, input schema, and output schema are fixed comparison controls. A core change requires an explicit team decision, a new version, and retesting of every accepted agent.")
add_bullet(doc, "Specialist layer  ", "Each owner may revise specialist instructions, cases, assets, responses, and records within the agent folder. Changes must not bypass the shared contract.")
add_bullet(doc, "Evidence layer  ", "Every candidate must preserve its baseline, first run, weakness or failure, revision, retest, validation result, limitation, owner, and version.")

add_heading(doc, "Agent Acceptance Standard")
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_layout(table, (1.25, 3.55, 2.0))
headers = ("Criterion", "Required evidence", "Acceptance test")
for idx, text in enumerate(headers):
    table.rows[0].cells[idx].text = text
    shade_cell(table.rows[0].cells[idx], NAVY)
set_repeat_table_header(table.rows[0])
rows = [
    ("Contract", "Valid input and output under the frozen shared schema", "Validator passes and required fields are meaningful"),
    ("Accuracy", "Claims traced to source data or supplied course evidence", "Independent recomputation or source check agrees"),
    ("Behavior", "Primary case plus a contrast or transfer case when available", "Output remains useful without hiding a material limitation"),
    ("Learning", "A retained failure or weakness, targeted revision, and retest", "Revision addresses the diagnosed problem without breaking the core"),
    ("Ownership", "Named owner, current version, contribution record, and evidence path", "Another member can locate and reproduce the review"),
]
for ridx, values in enumerate(rows, start=1):
    cells = table.add_row().cells
    for idx, value in enumerate(values):
        cells[idx].text = value
        if ridx % 2 == 0:
            shade_cell(cells[idx], PALE_GRAY)
set_table_layout(table, (1.25, 3.55, 2.0))
set_table_borders(table)
style_table_text(table)

add_heading(doc, "Records and Verification", before=6)
add_label_paragraph(doc, "Required records  ", "Maintain the inventory, decision log, version and contribution record, evidence register, test and failure log, integration issue log, workflow evidence, and final-project handoff. Preserve earlier versions after material changes.", after=3)
add_label_paragraph(doc, "AI and contribution rule  ", "AI may assist with instructions, code, charts, explanations, and records, but an owner must verify claims against source data or authoritative course materials. Credit is attached to observable design, testing, verification, diagnosis, revision, integration, or review work.", after=2)

doc.add_page_break()

add_heading(doc, "Common Comparison Workflow", before=0)
steps = [
    ("1  Register", "Add the candidate to the team inventory with owner, purpose, version, tests, and evidence path."),
    ("2  Freeze", "Run the frozen-core check and record any mismatch before evaluating the output."),
    ("3  Test", "Run the same primary question and agreed contrast or transfer test. Preserve inputs and raw outputs."),
    ("4  Review", "Check schema compliance, factual accuracy, decision usefulness, limitations, and consistency with course principles."),
    ("5  Revise", "Log the most consequential failure or weakness, change the specialist layer, and repeat the affected tests."),
    ("6  Decide", "Mark the agent Accepted, Revise, or Retired. Record the basis, dissent, unresolved risk, owner, and next action."),
]
workflow = doc.add_table(rows=0, cols=2)
workflow.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_layout(workflow, (1.15, 5.65))
for idx, (label, description) in enumerate(steps):
    cells = workflow.add_row().cells
    cells[0].text = label
    cells[1].text = description
    shade_cell(cells[0], NAVY)
    if idx % 2 == 1:
        shade_cell(cells[1], PALE_GRAY)
set_table_borders(workflow)
for row in workflow.rows:
    row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for cidx, cell in enumerate(row.cells):
        set_cell_margins(cell, top=80, bottom=80)
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            for run in p.runs:
                set_font(run, size=8.8, bold=(cidx == 0), color=(RGBColor(255, 255, 255) if cidx == 0 else BLACK))
set_table_layout(workflow, (1.15, 5.65))

add_heading(doc, "Integration Decision and Preserved Dissent")
add_body(doc, "Our initial system will use the frozen core as the interface between specialist agents. An accepted specialist enters the shared inventory only after the comparison workflow is complete. If an agent cannot be exercised during the workshop, we will document the blocker, owner, and next test date rather than treating it as accepted.")
add_label_paragraph(doc, "Team decision  ", "[Confirm in workshop] Adopt the acceptance standard and workflow above, with any agreed edits recorded in the decision log.")
add_label_paragraph(doc, "Dissent or alternative  ", "________________________________________________________________________________")
add_label_paragraph(doc, "Unresolved issue  ", "Confirm the shared platform, repository location, transfer-test convention, and exact submission method with the instructor.")

add_heading(doc, "Final Project Handoff")
add_body(doc, "The final-project package will carry forward accepted agents, the frozen contract, test cases, source register, decision history, and unresolved risks. Any later core change triggers a new shared version and regression test of all accepted agents. Retired candidates remain in the decision history so the team can explain why they were excluded.")

add_heading(doc, "Initialized Workspace Status")
add_bullet(doc, "Inventory  ", "The shared inventory is initialized with the CYF Chart Improvement Agent as a candidate; remaining member agents must be added and compared.")
add_bullet(doc, "Exercised evidence  ", "One complete baseline, first-run, revision, retest, validation, and frozen-core evidence trail is preserved in the supporting-evidence folder.")
add_bullet(doc, "Open handoff  ", "The decision log, integration issue log, contribution record, evidence register, and final-project handoff are initialized for completion during the workshop.")

signoff = doc.add_table(rows=1, cols=3)
signoff.alignment = WD_TABLE_ALIGNMENT.LEFT
labels = ("TEAM CONFIRMATION", "□ Adopted   □ Revised", "Date  __________________")
widths = (1.55, 3.55, 1.7)
set_table_layout(signoff, widths)
for idx, (label, width) in enumerate(zip(labels, widths)):
    cell = signoff.rows[0].cells[idx]
    cell.text = label
    if idx == 0:
        shade_cell(cell, PALE_BLUE)
set_table_borders(signoff)
style_table_text(signoff, header=False)
for idx, cell in enumerate(signoff.rows[0].cells):
    for p in cell.paragraphs:
        for run in p.runs:
            set_font(run, size=8.6, bold=(idx == 0))

doc.settings.odd_and_even_pages_header_footer = True
section.different_first_page_header_footer = False
for footer in (section.footer, section.even_page_footer):
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.space_before = Pt(0)
    footer_p.paragraph_format.space_after = Pt(0)
    set_font(footer_p.add_run("Team Workshop 2  |  Team Agent Design and Integration Standard"), size=8.2, color=GRAY)

core_props = doc.core_properties
core_props.title = "Team Agent Design and Integration Standard"
core_props.subject = "Team Workshop 2"
core_props.keywords = "agentic AI, agent design, integration, workflow evidence"

doc.save(OUTPUT)
print(OUTPUT)
